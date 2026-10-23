import io
import os
import streamlit as st
from google import genai
from google.genai import types

# ---------- Configuración (cambia estos textos a tu gusto) ----------
MODELO = "gemini-3.5-flash-lite"
TITULO = "🤖 Hola Soy Adrian, tu asistente de AI virtual. ¿En qué te puedo ayudar?"
INSTRUCCIONES = (
    "Eres un asistente amable y claro. Responde siempre en español. "
    "Cuando el usuario suba archivos o imágenes, analízalos con detalle, "
    "explica su contenido y desarrolla lo que te pida (resúmenes, explicaciones, "
    "correcciones, ideas, etc.)."
)

IMAGENES = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "webp": "image/webp"}
EXTENSIONES = ["png", "jpg", "jpeg", "webp", "pdf", "txt", "md", "csv", "docx"]

st.set_page_config(page_title="Mi Chatbot", page_icon="🤖")
st.title(TITULO)

api_key = os.getenv("GEMINI_API_KEY")
if not api_key:
    st.error("Falta la variable de entorno GEMINI_API_KEY.")
    st.stop()


# ---------- Funciones ----------
def archivo_a_parte(archivo):
    """Convierte un archivo subido en algo que Gemini pueda leer."""
    nombre = archivo.name
    ext = nombre.rsplit(".", 1)[-1].lower()
    datos = archivo.getvalue()

    if ext in IMAGENES:
        return types.Part.from_bytes(data=datos, mime_type=IMAGENES[ext])

    if ext == "pdf":
        return types.Part.from_bytes(data=datos, mime_type="application/pdf")

    if ext in ("txt", "md", "csv"):
        texto = datos.decode("utf-8", errors="ignore")
        return f"[Contenido del archivo {nombre}]\n{texto}"

    if ext == "docx":
        from docx import Document

        doc = Document(io.BytesIO(datos))
        partes = [p.text for p in doc.paragraphs]
        for tabla in doc.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(celda.text for celda in fila.cells))
        return f"[Contenido del archivo {nombre}]\n" + "\n".join(partes)

    return f"[No se pudo leer el archivo {nombre}]"


def mostrar_mensaje(m):
    with st.chat_message(m["rol"]):
        for img in m.get("imagenes", []):
            st.image(img, width=300)
        for nombre in m.get("adjuntos", []):
            st.caption(f"📎 {nombre}")
        if m["texto"]:
            st.markdown(m["texto"])


# ---------- Estado de la sesión ----------
if "client" not in st.session_state:
    try:
        st.session_state.client = genai.Client(
            api_key=api_key,
            http_options=types.HttpOptions(
                retry_options=types.HttpRetryOptions(attempts=5)
            ),
        )
    except Exception:
        st.session_state.client = genai.Client(api_key=api_key)

if "chat" not in st.session_state:
    st.session_state.chat = st.session_state.client.chats.create(
        model=MODELO,
        config=types.GenerateContentConfig(system_instruction=INSTRUCCIONES),
    )
    st.session_state.mensajes = []

if st.sidebar.button("🗑️ Nueva conversación"):
    del st.session_state["chat"]
    st.rerun()



# ---------- Mostrar historial ----------
for m in st.session_state.mensajes:
    mostrar_mensaje(m)

# ---------- Entrada del usuario ----------
entrada = st.chat_input(
    "Escribe tu mensaje o adjunta archivos con el clip 📎",
    accept_file="multiple",
    file_type=EXTENSIONES,
)

if entrada:
    texto = (entrada.text or "").strip()
    archivos = entrada.files or []

    if not texto and archivos:
        texto = "Analiza y explica detalladamente el contenido de lo que te envié."

    mensaje_usuario = {
        "rol": "user",
        "texto": texto,
        "imagenes": [
            a.getvalue() for a in archivos
            if a.name.rsplit(".", 1)[-1].lower() in IMAGENES
        ],
        "adjuntos": [
            a.name for a in archivos
            if a.name.rsplit(".", 1)[-1].lower() not in IMAGENES
        ],
    }
    st.session_state.mensajes.append(mensaje_usuario)
    mostrar_mensaje(mensaje_usuario)

    with st.chat_message("assistant"):
        try:
            contenido = [archivo_a_parte(a) for a in archivos]
            contenido.append(texto)

            def generar():
                for parte in st.session_state.chat.send_message_stream(contenido):
                    if parte.text:
                        yield parte.text

            respuesta = st.write_stream(generar())
        except Exception as e:
            respuesta = f"⚠️ Error: {e}"
            st.error(respuesta)

    st.session_state.mensajes.append({"rol": "assistant", "texto": respuesta})